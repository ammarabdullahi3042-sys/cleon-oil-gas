from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
from xml.sax.saxutils import escape
from pathlib import Path

root=Path('output/reports')
RED=colors.HexColor('#FF0000'); GREEN=colors.HexColor('#008000'); BLACK=colors.black
def clean(text): return escape(text).replace('’',"'").replace('“','"').replace('”','"').replace('–','-').replace('—','-')

for src in root.glob('*.docx'):
    d=Document(src); dest=src.with_suffix('.pdf')
    styles=getSampleStyleSheet()
    title=ParagraphStyle('SampleTitle',parent=styles['Title'],fontName='Times-Bold',fontSize=20,leading=24,textColor=RED,alignment=TA_CENTER,spaceBefore=12,spaceAfter=28)
    subtitle=ParagraphStyle('SampleSubtitle',parent=styles['Normal'],fontName='Times-Bold',fontSize=12,leading=15,textColor=GREEN,alignment=TA_CENTER,spaceAfter=2)
    official=ParagraphStyle('Official',parent=subtitle,fontSize=11,leading=14,spaceAfter=14)
    meta=ParagraphStyle('Meta',parent=styles['Normal'],fontName='Times-Roman',fontSize=11,leading=14,textColor=BLACK,alignment=TA_CENTER,spaceAfter=44)
    body=ParagraphStyle('SampleBody',parent=styles['BodyText'],fontName='Times-Roman',fontSize=11,leading=15,textColor=BLACK,alignment=TA_JUSTIFY,spaceAfter=5)
    section=ParagraphStyle('SampleSection',parent=body,fontName='Helvetica-Bold',fontSize=12,leading=15,textColor=RED,alignment=0,spaceBefore=9,spaceAfter=4,keepWithNext=True)
    attendance_section=ParagraphStyle('AttendanceSection',parent=section,keepWithNext=False)
    green=ParagraphStyle('GreenLabel',parent=body,fontName='Times-Roman',textColor=GREEN,spaceAfter=4)
    small=ParagraphStyle('Attendance',parent=body,fontName='Times-Roman',fontSize=10.5,leading=13,alignment=0,spaceAfter=0)
    story=[]
    title_text=d.paragraphs[0].text.upper()
    subtitle_text=d.paragraphs[1].text.upper()
    story.append(Paragraph(clean(title_text),title))
    story.append(Paragraph(clean(subtitle_text),subtitle))
    story.append(Paragraph('OFFICIAL MINUTES',official))
    meta_parts=[f'{row.cells[0].text}: {row.cells[1].text}' for row in d.tables[0].rows if row.cells[1].text.strip()]
    meta_line=' | '.join(meta_parts)
    story.append(Paragraph(clean(meta_line),meta))
    paras=d.paragraphs[3:]
    # Put the opening and attendee register at the front, as in the supplied minutes.
    opening=[]; rest=[]; attendance_started=False
    for p in paras:
        txt=p.text.strip()
        if not txt or txt.startswith('Names recorded:') or txt.startswith('Coordinator') or txt.startswith('_') or txt.startswith('Hajiya Khadija Gambo Hawaja'): continue
        if txt == 'Attendance':
            attendance_started=True; continue
        is_head=p.runs and all((not r.text.strip() or r.bold) for r in p.runs)
        if not attendance_started and not opening and is_head:
            opening.append(Paragraph(clean(txt),section))
        elif not attendance_started and opening and not is_head:
            opening.append(Paragraph(clean(txt),body)); attendance_started=True
        elif not attendance_started:
            rest.append((is_head,txt))
        else:
            rest.append((is_head,txt))
    story.extend(opening)
    story.append(Paragraph('Full Attendance List',attendance_section))
    attendance=d.tables[-1]
    rows=[[Paragraph('<b>S/No.</b>',small),Paragraph('<b>Name</b>',small),Paragraph('<b>Phone No.</b>',small)]]
    for row in attendance.rows[1:]: rows.append([Paragraph(clean(row.cells[i].text),small) for i in range(3)])
    tab=Table(rows,colWidths=[.55*inch,3.85*inch,2.25*inch],repeatRows=1,hAlign='LEFT')
    tab.setStyle(TableStyle([('LINEABOVE',(0,0),(-1,0),.6,BLACK),('LINEBELOW',(0,0),(-1,0),.6,BLACK),('LINEBELOW',(0,-1),(-1,-1),.6,BLACK),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5)]))
    story.append(tab); story.append(Spacer(1,12))
    for is_head,txt in rest:
        story.append(Paragraph(clean(txt),section if is_head else body))
    signature_name=Paragraph('Hajiya Khadija Gambo Hawaja',body)
    signature_role=Paragraph('Coordinator',body)
    signature_block=KeepTogether([Spacer(1,14),HRFlowable(width=2.2*inch,thickness=.8,color=BLACK,hAlign='LEFT',spaceBefore=0,spaceAfter=6),signature_name,signature_role])
    story.append(signature_block)
    pdf=SimpleDocTemplate(str(dest),pagesize=letter,rightMargin=.45*inch,leftMargin=.45*inch,topMargin=.38*inch,bottomMargin=.48*inch,title=d.core_properties.title or src.stem,author='')
    pdf.build(story)
    print(dest)
