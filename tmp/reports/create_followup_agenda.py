from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT=Path('output/reports/Jos_North_Data_Exercise_Follow_Up_Agenda_30_October_2026.docx')
d=Document(); sec=d.sections[0]
sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.6); sec.right_margin=Inches(.6)
normal=d.styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(10.5); normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.0

RED=RGBColor(255,0,0); GREEN=RGBColor(0,128,0)
title=d.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before=Pt(10); title.paragraph_format.space_after=Pt(18)
r=title.add_run('JOS NORTH WARD CHAIRMEN FOLLOW-UP MEETING'); r.font.name='Times New Roman'; r.font.size=Pt(19); r.font.bold=True; r.font.color.rgb=RED
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(2)
r=p.add_run('POLLING-UNIT DATA AND PASSPORT PHOTOGRAPH REQUIREMENTS'); r.font.name='Times New Roman'; r.font.size=Pt(11.5); r.font.bold=True; r.font.color.rgb=GREEN
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(11)
r=p.add_run('MEETING AGENDA AND MINUTES SHEET'); r.font.name='Times New Roman'; r.font.size=Pt(11); r.font.bold=True; r.font.color.rgb=GREEN
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(20)
r=p.add_run('Date: 30 October 2026  |  Venue: PDP Jos North Party Office  |  Time: __________________'); r.font.name='Times New Roman'; r.font.size=Pt(10.5)

def shade(c,color='E8F2EC'):
    pr=c._tc.get_or_add_tcPr(); el=OxmlElement('w:shd'); el.set(qn('w:fill'),color); pr.append(el)
def borders(t):
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single' if e in ('top','bottom','insideH') else 'nil'); x.set(qn('w:sz'),'5'); x.set(qn('w:color'),'000000'); b.append(x)
    t._tbl.tblPr.append(b)
def table(headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths): t.columns[i].width=Inches(w)
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c)
        for rr in c.paragraphs[0].runs: rr.font.name='Times New Roman'; rr.font.size=Pt(9); rr.font.bold=True
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):
            cells[i].text=str(val); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
                for rr in p.runs: rr.font.name='Times New Roman'; rr.font.size=Pt(9)
    borders(t); return t
def heading(text):
    p=d.add_paragraph(); p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(text); r.font.name='Arial'; r.font.size=Pt(11.5); r.font.bold=True; r.font.color.rgb=RED
def para(text):
    p=d.add_paragraph(text); p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; return p

heading('Background')
para('This meeting continues the discussions held on 28 and 29 October 2026 concerning the proposed process for compiling data on 120 persons from each polling unit. At the 28 October consultation, ward chairmen raised concerns about logistics, stakeholder inclusion, trust, agreed facilitation and community resistance. The follow-up was arranged to broaden consultation and develop practical recommendations ahead of the elections.')
para('A further concern for today’s discussion is that many people cannot afford the ₦1,000 required for passport photographs. The chairmen also have limited means to pay this cost on behalf of residents. The meeting is to consider a workable and fair response before the data exercise proceeds.')

heading('Meeting Objectives')
for x in [
'Clarify whether a passport photograph is mandatory for each person and confirm the exact specification accepted by the campaign data team.',
'Assess the cost barrier created by the ₦1,000 photograph fee and its effect on participation across polling units.',
'Consider practical options for reducing or meeting the cost without placing an unaffordable burden on residents or ward chairmen.',
'Confirm logistics, responsibilities, stakeholder consultation and the process for reporting community feedback.'
]:
    p=d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(x)

heading('Proposed Agenda')
agenda=[
('1','Opening prayer, welcome and confirmation of attendance',''),
('2','Review of concerns and actions arising from the 28–29 October discussions',''),
('3','Clarification of the 120-person-per-polling-unit data process and photograph requirement',''),
('4','Discussion of the ₦1,000 passport photograph cost and affordability barrier',''),
('5','Consideration of support options: campaign-funded photographs, a central photo day, a subsidy/waiver, or an approved alternative format',''),
('6','Logistics, funding responsibility, payment timing and ward-level communication',''),
('7','Consultation with Christian community ward chairmen and preparation of a consolidated report',''),
('8','Agreement on decisions, responsible persons and deadlines; closing remarks and prayer','')]
table(['No.','Agenda item','Notes / decision'],agenda,[.45,4.5,2.1])

heading('Options for Discussion')
options=[
('Campaign-supported photo session','Arrange a central or ward-based photo day and cover the agreed cost from campaign logistics funds.'),
('Subsidy or waiver','Identify residents who cannot pay and agree a transparent support or waiver process.'),
('Alternative format','Ask the campaign data team whether an existing digital photograph or another approved identification method can be accepted.'),
('Staged collection','Do not exclude residents solely for inability to pay until the requirement, support method and funding are confirmed.')]
table(['Option','Points to consider'],options,[2.0,5.05])
p=d.add_paragraph(); p.paragraph_format.space_before=Pt(2)
r=p.add_run('These are proposals for the meeting to consider; no option is recorded as approved.'); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(80,80,80)

heading('Attendance')
att=[
('Hajiya Khadijah Gambo Hawaja','Meeting convenor',''),
('Barr. Aminu Sani','Coordinator, Jos North',''),
('Jos North PDP Chairman','Chairman',''),
('Jos North PDP Secretary','Secretary',''),
('Musa Lawal Ahmad','Gangare Ward Chairman','08035015762'),
('Alh. Yawale Falke','Garba Daho Ward Chairman','08033342449'),
('Kawu Babanjoda','Naraguta A Ward Chairman','07035586406'),
('Ado Garba Jibiya','Abba na Shehu Ward Chairman','07038746261'),
('Musa Yunusa Ubale Ibrahim','Katsina Ward Chairman','07067985449'),
('Alkasim Musa Wetka','Sarkin Arab Ward Chairman','08063970491'),
('Alhassan Umar Aliyu Ali Kazaure','Kazaure Ward Chairman','07060905093'),
('Christian community ward chairmen','Names to be recorded at meeting','')]
table(['Name','Role / Ward','Phone No.'],att,[2.6,3.3,1.15])
p=d.add_paragraph('Please mark actual attendance and amend the register during the meeting.'); p.paragraph_format.space_before=Pt(2)
r=p.runs[0]; r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(80,80,80)

heading('Decisions and Action Points')
table(['Decision / action','Responsible person','Due date / status'],[('','',''),('','',''),('','',''),('','','')],[3.2,2.1,1.75])
heading('Closing')
para('The meeting is expected to conclude with clear guidance on the photograph requirement, how the cost barrier will be addressed, and who will coordinate the next steps with ward leaders and the campaign data team.')
p=d.add_paragraph(); p.paragraph_format.space_before=Pt(15); p.paragraph_format.space_after=Pt(4)
r=p.add_run('________________________________'); r.font.name='Times New Roman'; r.font.size=Pt(10)
p=d.add_paragraph(); r=p.add_run('Hajiya Khadijah Gambo Hawaja\nMeeting Convenor'); r.font.name='Times New Roman'; r.font.size=Pt(10)
p=d.add_paragraph(); p.paragraph_format.space_before=Pt(10); r=p.add_run('Secretary: ______________________________    Signature: ______________________________'); r.font.name='Times New Roman'; r.font.size=Pt(10)

d.core_properties.title='Jos North Ward Chairmen Follow-Up Meeting Agenda and Minutes Sheet'
d.core_properties.subject='Passport photograph affordability and polling-unit data process follow-up'
d.save(OUT); print(OUT)
