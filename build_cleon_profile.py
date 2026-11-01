from pathlib import Path
from shutil import copy2
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).parent/'CLEON_Company_Profile'
ASSET=ROOT/'assets'
ASSET.mkdir(parents=True,exist_ok=True)
src=Path(r'C:\Users\ammar\Downloads\cleon-logo-for-dark-backgrounds.png')
dark_logo=ASSET/'cleon-logo-dark-background.png'
white_logo=ASSET/'cleon-logo-white-background.png'
copy2(src,dark_logo)
im=Image.open(src).convert('RGBA')
px=im.load()
for y in range(im.height):
    for x in range(im.width):
        r,g,b,a=px[x,y]
        if a>0 and r>220 and g>220 and b>220:
            px[x,y]=(19,27,38,a)
im.save(white_logo)
(ASSET/'Asset_Notes.txt').write_text(
    'CLEON Company Profile assets\n\n'
    'cleon-logo-dark-background.png: source logo supplied for dark surfaces.\n'
    'cleon-logo-white-background.png: retained optional light-surface derivative.\n'
    'oil-and-gas-hero-source.png: original user-supplied photograph.\n'
    'oil-and-gas-hero-dark-blend.png and oil-and-gas-hero-white-blend.png: soft-edge blends for profile covers.\n'
    'The profile source documents are editable Microsoft Word DOCX files.\n',
    encoding='utf-8')

INK='F3F5F7'; MUTED='B7C0CA'; WHITE='FFFFFF'; NAVY='131B26'; NAVY2='1B2734'; SLATE='233344'
ORANGE='F47A20'; GOLD='DDAE67'; PALE='EDF1F5'; PALE2='F6F8FA'; LINE_D='405164'; LINE_L='D6DEE6'
BODY_FONT='Aptos'; HEAD_FONT='Aptos Display'

def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=tcPr.find(qn('w:shd'))
    if shd is None:
        shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),color)

def cell_margins(cell,top=100,start=130,bottom=100,end=130):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr()
    mar=tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar=OxmlElement('w:tcMar'); tcPr.append(mar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=mar.find(qn('w:'+m))
        if node is None: node=OxmlElement('w:'+m); mar.append(node)
        node.set(qn('w:w'),str(v)); node.set(qn('w:type'),'dxa')

def table_borders(table,color):
    tblPr=table._tbl.tblPr
    b=tblPr.first_child_found_in('w:tblBorders')
    if b is None: b=OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag=qn('w:'+edge); el=b.find(tag)
        if el is None: el=OxmlElement('w:'+edge); b.append(el)
        el.set(qn('w:val'),'single' if edge in ('bottom','insideH') else 'nil')
        el.set(qn('w:sz'),'5'); el.set(qn('w:space'),'0'); el.set(qn('w:color'),color)

def set_repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def set_run(run,size=11,bold=False,color=None,font=BODY_FONT,italic=False):
    run.font.name=font
    run._element.rPr.rFonts.set(qn('w:eastAsia'),font)
    run.font.size=Pt(size); run.bold=bold; run.italic=italic
    if color: run.font.color.rgb=RGBColor.from_string(color)
    return run

def set_keep(p,keep_next=False):
    pPr=p._p.get_or_add_pPr()
    if keep_next:
        el=OxmlElement('w:keepNext'); pPr.append(el)
    return p

def add_rule(p,color):
    pPr=p._p.get_or_add_pPr()
    pBdr=pPr.find(qn('w:pBdr'))
    if pBdr is None: pBdr=OxmlElement('w:pBdr'); pPr.append(pBdr)
    bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'9'); bottom.set(qn('w:space'),'7'); bottom.set(qn('w:color'),color); pBdr.append(bottom)

def set_bg(doc,color):
    bg=OxmlElement('w:background'); bg.set(qn('w:color'),color)
    doc._element.insert(0,bg)
    settings=doc.settings._element
    compat=settings.find(qn('w:displayBackgroundShape'))
    if compat is None:
        compat=OxmlElement('w:displayBackgroundShape'); settings.append(compat)

def page_num(p,color):
    r=p.add_run()
    set_run(r,9,color=color)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE')
    inner=OxmlElement('w:r'); pr=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),color); pr.append(c)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'18'); pr.append(sz)
    t=OxmlElement('w:t'); t.text='2'; inner.append(pr); inner.append(t); fld.append(inner); p._p.append(fld)

def add_para(doc,text='',size=11,color=None,bold=False,space_after=6,align=None,italic=False):
    p=doc.add_paragraph()
    if align is not None: p.alignment=align
    p.paragraph_format.space_after=Pt(space_after)
    p.paragraph_format.line_spacing=1.12
    if text: set_run(p.add_run(text),size,bold,color,italic=italic)
    return p

def add_heading(doc,text,dark,level=1):
    c=GOLD if dark else '995416'
    size=25 if level==1 else 15
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(3 if level==1 else 12)
    p.paragraph_format.space_after=Pt(8 if level==1 else 4)
    set_keep(p,True); set_run(p.add_run(text),size,True,c,HEAD_FONT)
    if level==1: add_rule(p,ORANGE if dark else 'D9A46A')
    return p

def add_kicker(doc,text,dark):
    p=doc.add_paragraph()
    p.paragraph_format.space_before=Pt(0); p.paragraph_format.space_after=Pt(5)
    set_run(p.add_run(text.upper()),9,True,ORANGE if dark else 'B65612')
    return p

def add_bullet(doc,lead,body,dark):
    p=doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent=Inches(.24); p.paragraph_format.first_line_indent=Inches(-.15)
    p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.1
    col=INK if dark else NAVY
    set_run(p.add_run(lead+' '),11,True,GOLD if dark else '995416')
    set_run(p.add_run(body),11,False,col)
    return p

def add_facts(doc,rows,dark):
    t=doc.add_table(rows=0,cols=2)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    t.columns[0].width=Inches(1.75); t.columns[1].width=Inches(4.9)
    table_borders(t,LINE_D if dark else LINE_L)
    for i,(k,v) in enumerate(rows):
        cells=t.add_row().cells
        cells[0].width=Inches(1.75); cells[1].width=Inches(4.9)
        for c in cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(c)
            shade(c,(NAVY2 if i%2==0 else '18232F') if dark else (PALE2 if i%2==0 else WHITE))
        set_run(cells[0].paragraphs[0].add_run(k),10,True,GOLD if dark else '995416')
        set_run(cells[1].paragraphs[0].add_run(v),10.5,False,INK if dark else NAVY)
        for p in (cells[0].paragraphs[0],cells[1].paragraphs[0]):
            p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
    return t

def setup(dark):
    d=Document()
    sec=d.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.62)
    sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
    sec.header_distance=Inches(.32); sec.footer_distance=Inches(.35)
    sec.different_first_page_header_footer=True
    set_bg(d,NAVY if dark else WHITE)
    normal=d.styles['Normal']; normal.font.name=BODY_FONT; normal.font.size=Pt(11); normal.font.color.rgb=RGBColor.from_string(INK if dark else NAVY)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'),BODY_FONT)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.12
    bullet=d.styles['List Bullet']; bullet.font.name=BODY_FONT; bullet.font.size=Pt(11); bullet.font.color.rgb=RGBColor.from_string(INK if dark else NAVY)
    header=sec.header.paragraphs[0]
    header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_run(header.add_run('CLEON  /  COMPANY PROFILE'),8,True,MUTED if dark else '637180')
    add_rule(header,LINE_D if dark else LINE_L)
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.LEFT
    set_run(footer.add_run('CLEON OIL & GAS LTD    •    NIGERIA'),8,True,MUTED if dark else '637180')
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.55))
    set_run(footer.add_run('\t'),8,color=MUTED if dark else '637180'); page_num(footer,MUTED if dark else '637180')
    return d

def pagebreak(d):
    d.add_page_break()

def build(dark):
    d=setup(dark)
    base=INK if dark else NAVY
    secondary=MUTED if dark else '586676'
    accent=GOLD if dark else '995416'
    # Cover
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(5)
    if dark:
        p.add_run().add_picture(str(dark_logo),width=Inches(2.45))
    else:
        plaque=d.add_table(rows=1,cols=1)
        plaque.alignment=WD_TABLE_ALIGNMENT.LEFT; plaque.autofit=False
        plaque.columns[0].width=Inches(2.62)
        cell=plaque.cell(0,0); cell.width=Inches(2.62); shade(cell,NAVY)
        cell_margins(cell,top=80,start=100,bottom=80,end=100)
        table_borders(plaque,NAVY)
        lp=cell.paragraphs[0]; lp.paragraph_format.space_after=Pt(0)
        lp.add_run().add_picture(str(dark_logo),width=Inches(2.45))
    add_kicker(d,'COMPANY PROFILE  /  INVESTOR OVERVIEW',dark)
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.0
    set_run(p.add_run('CLEON OIL & GAS LTD'),29,True,base,HEAD_FONT)
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(8)
    set_run(p.add_run('Flexible Energy. Reliable Service.'),16,True,accent,HEAD_FONT)
    add_rule(p,ORANGE if dark else 'D9A46A')
    ph=d.add_paragraph(); ph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    ph.paragraph_format.space_before=Pt(5); ph.paragraph_format.space_after=Pt(5)
    ph.add_run().add_picture(str(ASSET/('oil-and-gas-hero-dark-blend.png' if dark else 'oil-and-gas-hero-white-blend.png')),width=Inches(6.0))
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.12
    set_run(p.add_run('A Nigerian downstream energy company focused on disciplined fuel retail, forecourt management and technology-supported operations.'),13,False,base)
    add_para(d,'PMS & AGO RETAIL   /   STATION MANAGEMENT   /   NON-FUEL RETAIL   /   CLEAN-ENERGY OPPORTUNITIES',8.5,secondary,True,space_after=12)
    add_para(d,'REGISTERED IN NIGERIA  •  RC 8817624',9.5,accent,True,space_after=1)
    add_para(d,'Investor and partner overview  |  2026',8.5,secondary,space_after=0)
    # Company
    pagebreak(d)
    add_kicker(d,'01  /  COMPANY',dark); add_heading(d,'A focused downstream operator',dark)
    add_para(d,'CLEON Oil & Gas Ltd is a Nigerian downstream energy company focused on fuel retail, station management, non-fuel retail development and technology-enabled forecourt services.',12,base,space_after=10)
    add_para(d,'The business plan identifies a pilot station in Kafanchan, Kaduna State, with operations beginning in October 2025. CLEON’s operating approach pairs day-to-day station controls with measured site and service expansion.',11,secondary,space_after=12)
    add_heading(d,'Company at a glance',dark,2)
    add_facts(d,[
        ('Legal name','CLEON Oil & Gas Ltd'),
        ('Registration','RC 8817624, Nigeria'),
        ('Sector','Downstream oil and gas retail'),
        ('Pilot location','Kafanchan, Kaduna State'),
        ('Operating scope','PMS and AGO retail, station management, non-fuel retail'),
    ],dark)
    add_heading(d,'Vision',dark,2)
    add_para(d,'To set the standard for accurate, safe and transparent fuel retail across Nigeria.',14,accent,True,space_after=12)
    add_heading(d,'Mission',dark,2)
    add_para(d,'To deliver high-quality white products (PMS and AGO) and clean-energy solutions (LPG and CNG) through an optimized network of leased and managed retail forecourts.',11,base,space_after=8)
    # Services and model
    pagebreak(d)
    add_kicker(d,'02  /  SERVICES & OPERATING MODEL',dark); add_heading(d,'Forecourt services built around control',dark)
    add_para(d,'CLEON’s business model combines fuel retail with station operating services and adjacent forecourt opportunities. The mix is designed to fit the needs and readiness of each site.',12,base,space_after=10)
    add_bullet(d,'PMS & AGO retail.','Retail operations supported by receipt, pump-sales and tank-dip records.',dark)
    add_bullet(d,'Station management.','Operating support for CLEON-operated and third-party forecourts, with clear reconciliation and reporting processes.',dark)
    add_bullet(d,'Non-fuel retail.','Development of complementary retail services and commercial opportunities around the forecourt.',dark)
    add_bullet(d,'LPG & CNG.','Clean-energy options identified for phased evaluation and future expansion.',dark)
    add_heading(d,'A flexible operating structure',dark,2)
    add_para(d,'CLEON’s plan is based on leased and managed retail sites. This allows the company to evaluate each location, supply route, operating partner and investment requirement before committing to an expansion step.',11,secondary,space_after=12)
    add_facts(d,[
        ('Site readiness','Review equipment, utilities, access, local competition and safety needs.'),
        ('Operating controls','Set product, cash, stock, reconciliation and reporting routines.'),
        ('Partner model','Define operating responsibilities and information flow for managed sites.'),
        ('Expansion','Assess non-fuel retail and clean-energy opportunities as site conditions allow.'),
    ],dark)
    # Technology
    pagebreak(d)
    add_kicker(d,'03  /  TECHNOLOGY & REPORTING',dark); add_heading(d,'Operational visibility from the forecourt',dark)
    add_para(d,'CLEON has developed a station operations application for key forecourt records. It brings daily operating information into a more consistent view for station teams and management.',12,base,space_after=10)
    add_bullet(d,'Stock and sales.','Records receipts, pump sales, tank dips and stock positions to support variance review.',dark)
    add_bullet(d,'Commercial control.','Tracks selling prices and product costs to support gross-profit and net-profit reporting.',dark)
    add_bullet(d,'Payment visibility.','Records cash, POS, deposits, transfers and account balances for reconciliation.',dark)
    add_bullet(d,'Management reports.','Supports Excel, PDF and CSV exports for internal and owner reporting.',dark)
    add_heading(d,'Website and investor information',dark,2)
    add_para(d,'CLEON’s website is already created and provides the company’s online presence. CLEON intends to use it to give authorized investors a clear view of contributions, deployment milestones and available reports, with timely progress updates.',11,base,space_after=8)
    add_para(d,'The current station application stores desktop data locally. Live investor balances, automated alerts and connected reporting should be enabled only after secure hosting, access controls, data synchronization, audit history and backups are in place.',10.5,secondary,space_after=10)
    add_heading(d,'Technology roadmap',dark,2)
    add_facts(d,[
        ('In use','Local desktop application for pilot station records and report exports.'),
        ('Next','Regular backups, user roles and an audit trail for edits and approvals.'),
        ('Later','Secure shared service with remote multi-site access and investor views.'),
    ],dark)
    # Growth and investor
    pagebreak(d)
    add_kicker(d,'04  /  GROWTH & INVESTOR ENGAGEMENT',dark); add_heading(d,'Scale in stages, report with clarity',dark)
    add_para(d,'CLEON’s expansion path is intended to follow operating readiness: stabilize controls at the pilot, review opportunities with site-level evidence, and expand services when supply, safety, people and funding requirements are understood.',12,base,space_after=10)
    add_heading(d,'Priority workstreams',dark,2)
    add_bullet(d,'Strengthen the pilot.','Build reliable stock, sales, payment and maintenance records and close routine reconciliations.',dark)
    add_bullet(d,'Prepare for replication.','Confirm backup procedures, user permissions, audit history and consistent reporting before remote multi-site operation.',dark)
    add_bullet(d,'Evaluate growth opportunities.','Assess additional managed or leased sites, non-fuel retail and future LPG/CNG opportunities against local conditions.',dark)
    add_heading(d,'Investor visibility',dark,2)
    add_para(d,'The planned investor information area is intended to make the use and progress of investment easier to follow. Subject to secure implementation and verified records, it can present:',11,secondary,space_after=7)
    add_bullet(d,'Capital activity.','Contributions received and their recorded deployment.',dark)
    add_bullet(d,'Milestone progress.','Status of agreed site, equipment and operating milestones.',dark)
    add_bullet(d,'Business updates.','Available operating and management reports, with timely notices of material progress.',dark)
    add_para(d,'CLEON should distinguish actual results from targets and estimates in investor communications. The content of each update should be based on reconciled business records.',10.5,secondary,space_after=10)
    add_heading(d,'Company details',dark,2)
    add_facts(d,[
        ('Company','CLEON Oil & Gas Ltd'),
        ('Registration','RC 8817624'),
        ('Country','Nigeria'),
        ('Head office and contact','Details to be confirmed by company management.'),
    ],dark)
    out=ROOT/f'CLEON_Company_Profile_{"Dark" if dark else "White"}.docx'
    d.save(out)
    return out

for dark in (True,False):
    print(build(dark))

