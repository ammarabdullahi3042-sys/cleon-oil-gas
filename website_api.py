"""Small authenticated API for publishing app snapshots to a hosted website."""

from contextlib import contextmanager
import csv
import hashlib
import hmac
from html import escape
import io
import json
import mimetypes
import os
import re
import secrets
import sqlite3
import sys
import threading
import time
from datetime import datetime, timezone
from getpass import getpass
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse


MAX_BODY_BYTES = 5 * 1024 * 1024
MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
SESSION_HOURS = 8
PASSWORD_ITERATIONS = 240_000
SECTIONS = {"dashboard", "reports"}
ROLES = {"super_admin", "admin", "manager", "member"}
ROLE_ACCESS = {
    "super_admin": {"members", "dashboard", "reports", "profile"},
    "admin": {"members", "dashboard", "reports"},
    "manager": {"dashboard", "reports"},
    "member": {"reports", "profile"},
}
PREVIEW_TYPES = {
    "application/pdf": b"%PDF-",
    "image/jpeg": b"\xff\xd8\xff",
    "image/png": b"\x89PNG\r\n\x1a\n",
    "image/gif": b"GIF8",
    "image/webp": b"RIFF",
}
LOGIN_ATTEMPTS = {}
LOGIN_ATTEMPTS_LOCK = threading.Lock()


def _database_path():
    return Path(os.environ.get(
        "CLEON_WEBSITE_API_DATABASE",
        str(Path(__file__).resolve().with_name("website_snapshots.db")),
    ))


@contextmanager
def _connect_database():
    conn = sqlite3.connect(_database_path(), timeout=10)
    try:
        with conn:
            yield conn
    finally:
        conn.close()


def _initialize_database():
    with _connect_database() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS website_snapshots (
                section TEXT PRIMARY KEY,
                payload TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                published_at TEXT NOT NULL DEFAULT ''
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS website_members (
                email TEXT PRIMARY KEY,
                salt TEXT NOT NULL,
                password_hash TEXT NOT NULL,
                created_at TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'member'
            )
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS website_member_sessions (
                token_hash TEXT PRIMARY KEY,
                email TEXT NOT NULL REFERENCES website_members(email),
                expires_at TEXT NOT NULL
            )
        """)
        conn.execute("""
            CREATE INDEX IF NOT EXISTS idx_website_sessions_expiry
            ON website_member_sessions(expires_at)
        """)
        conn.execute("""
            CREATE TABLE IF NOT EXISTS website_documents (
                document_id TEXT PRIMARY KEY,
                filename TEXT NOT NULL,
                content_type TEXT NOT NULL,
                file_data BLOB NOT NULL,
                uploaded_at TEXT NOT NULL
            )
        """)
        member_columns = {
            row[1] for row in conn.execute(
                "PRAGMA table_info(website_members)"
            )
        }
        if "role" not in member_columns:
            conn.execute(
                "ALTER TABLE website_members ADD COLUMN "
                "role TEXT NOT NULL DEFAULT 'member'"
            )
        columns = {
            row[1] for row in conn.execute(
                "PRAGMA table_info(website_snapshots)"
            )
        }
        if "published_at" not in columns:
            conn.execute(
                "ALTER TABLE website_snapshots ADD COLUMN "
                "published_at TEXT NOT NULL DEFAULT ''"
            )


def _password_hash(password, salt):
    return hashlib.pbkdf2_hmac(
        "sha256", password.encode("utf-8"), bytes.fromhex(salt),
        PASSWORD_ITERATIONS,
    ).hex()


def _token_hash(token):
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


def _validate_role(role):
    if role not in ROLES:
        raise ValueError(f"Role must be one of: {', '.join(sorted(ROLES))}.")
    return role


def _add_member(email, role="member"):
    email = email.strip().casefold()
    if "@" not in email or len(email) > 254:
        raise ValueError("Enter a valid member email address.")
    role = _validate_role(role)
    password = getpass("New member password (at least 12 characters): ")
    if len(password) < 12:
        raise ValueError("Member passwords must be at least 12 characters.")
    confirmation = getpass("Confirm password: ")
    if not hmac.compare_digest(password, confirmation):
        raise ValueError("Passwords do not match.")
    salt = secrets.token_hex(16)
    with _connect_database() as conn:
        conn.execute(
            "INSERT INTO website_members(email,salt,password_hash,created_at,role) "
            "VALUES (?,?,?,?,?)",
            (
                email, salt, _password_hash(password, salt),
                datetime.now(timezone.utc).isoformat(timespec="seconds"),
                role,
            ),
        )
    print(f"Member account created for {email} with role {role}.")


def _set_member_role(email, role):
    email = email.strip().casefold()
    role = _validate_role(role)
    with _connect_database() as conn:
        cursor = conn.execute(
            "UPDATE website_members SET role=? WHERE email=?",
            (role, email),
        )
    if cursor.rowcount == 0:
        raise ValueError(f"No member account exists for {email}.")
    print(f"Role for {email} set to {role}.")


def _list_members():
    with _connect_database() as conn:
        rows = conn.execute(
            "SELECT email,role,created_at FROM website_members ORDER BY email"
        ).fetchall()
    for email, role, created_at in rows:
        print(f"{email}\t{role}\t{created_at}")


def _remove_member(email):
    email = email.strip().casefold()
    with _connect_database() as conn:
        conn.execute(
            "DELETE FROM website_member_sessions WHERE email=?", (email,)
        )
        cursor = conn.execute(
            "DELETE FROM website_members WHERE email=?", (email,)
        )
    if cursor.rowcount == 0:
        raise ValueError(f"No member account exists for {email}.")
    print(f"Member account removed for {email}; active sessions revoked.")


class WebsiteAPIHandler(BaseHTTPRequestHandler):
    server_version = "CLEONWebsiteAPI/1.0"

    def _send_json(self, status, payload):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self._add_cors_headers()
        self.end_headers()
        self.wfile.write(body)

    def _send_file(self, body, content_type, filename):
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header(
            "Content-Disposition", f'attachment; filename="{filename}"'
        )
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self._add_cors_headers()
        self.end_headers()
        self.wfile.write(body)

    def _send_document(self, body, content_type, filename, download=False):
        disposition = "attachment" if download else "inline"
        safe_ascii_name = filename.encode("ascii", "ignore").decode("ascii")
        safe_ascii_name = re.sub(r'[^A-Za-z0-9._ -]', "_", safe_ascii_name)
        encoded_name = quote(filename, safe="")
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header(
            "Content-Disposition",
            f'{disposition}; filename="{safe_ascii_name or "document"}"; '
            f"filename*=UTF-8''{encoded_name}",
        )
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self._add_cors_headers()
        self.end_headers()
        self.wfile.write(body)

    def _add_cors_headers(self):
        origin = self.headers.get("Origin")
        allowed_origin = os.environ.get(
            "CLEON_WEBSITE_ALLOWED_ORIGIN", ""
        ).strip().rstrip("/")
        if origin and allowed_origin and hmac.compare_digest(
            origin.rstrip("/"), allowed_origin
        ):
            self.send_header("Access-Control-Allow-Origin", allowed_origin)
            self.send_header("Vary", "Origin")
            self.send_header(
                "Access-Control-Allow-Headers",
                "Authorization, Content-Type, X-Document-Name",
            )
            self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")

    def do_OPTIONS(self):
        self.send_response(204)
        self._add_cors_headers()
        self.end_headers()

    def _upload_document(self):
        publish_token = os.environ.get("CLEON_WEBSITE_PUBLISH_TOKEN", "")
        if not self._authorized(publish_token):
            self._send_json(401, {"error": "Admin publishing authorization required"})
            return
        filename = unquote(self.headers.get("X-Document-Name", "")).strip()
        filename = filename.replace("\\", "/").rsplit("/", 1)[-1]
        if (
            not filename or len(filename) > 200
            or any(ord(character) < 32 for character in filename)
        ):
            self._send_json(400, {"error": "A valid document filename is required"})
            return
        content_type = mimetypes.guess_type(filename)[0]
        allowed_types = set(PREVIEW_TYPES) | {
            "application/msword",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/vnd.ms-excel",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-powerpoint",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        }
        if content_type not in allowed_types:
            self._send_json(415, {"error": "Upload a PDF, image, or Office document"})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            self._send_json(400, {"error": "Invalid Content-Length"})
            return
        if length <= 0 or length > MAX_DOCUMENT_BYTES:
            self._send_json(413, {"error": "Document must be under 10 MB"})
            return
        file_data = self.rfile.read(length)
        if len(file_data) != length:
            self._send_json(400, {"error": "Document upload was incomplete"})
            return
        signature = PREVIEW_TYPES.get(content_type)
        if signature and not file_data.startswith(signature):
            self._send_json(415, {"error": "File content does not match its type"})
            return
        if content_type == "image/webp" and file_data[8:12] != b"WEBP":
            self._send_json(415, {"error": "File content does not match its type"})
            return
        document_id = secrets.token_urlsafe(18)
        uploaded_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
        try:
            with _connect_database() as conn:
                conn.execute(
                    "INSERT INTO website_documents"
                    "(document_id,filename,content_type,file_data,uploaded_at) "
                    "VALUES (?,?,?,?,?)",
                    (document_id, filename, content_type, file_data, uploaded_at),
                )
        except sqlite3.Error:
            self._send_json(500, {"error": "Could not store document"})
            return
        self._send_json(201, {
            "document_id": document_id,
            "filename": filename,
            "content_type": content_type,
            "uploaded_at": uploaded_at,
        })

    def _list_documents(self):
        try:
            with _connect_database() as conn:
                documents = [
                    {
                        "document_id": document_id,
                        "filename": filename,
                        "content_type": content_type,
                        "uploaded_at": uploaded_at,
                        "previewable": content_type in PREVIEW_TYPES,
                    }
                    for document_id, filename, content_type, uploaded_at
                    in conn.execute(
                        "SELECT document_id,filename,content_type,uploaded_at "
                        "FROM website_documents ORDER BY uploaded_at DESC"
                    )
                ]
        except sqlite3.Error:
            self._send_json(500, {"error": "Could not list documents"})
            return
        self._send_json(200, {"documents": documents})

    def _get_document(self, parsed_path, member):
        document_id = unquote(parsed_path.path.removeprefix("/api/v1/documents/"))
        if not document_id or "/" in document_id or "\\" in document_id:
            self._send_json(404, {"error": "Document not found"})
            return
        download = (parse_qs(parsed_path.query).get("download") or [""])[0] == "1"
        if download and member["role"] not in ("admin", "super_admin"):
            self._send_json(403, {"error": "Document downloads require an admin role"})
            return
        try:
            with _connect_database() as conn:
                document = conn.execute(
                    "SELECT filename,content_type,file_data FROM website_documents "
                    "WHERE document_id=?",
                    (document_id,),
                ).fetchone()
        except sqlite3.Error:
            self._send_json(500, {"error": "Could not read document"})
            return
        if document is None:
            self._send_json(404, {"error": "Document not found"})
            return
        filename, content_type, file_data = document
        if not download and content_type not in PREVIEW_TYPES:
            self._send_json(415, {"error": "This document type cannot be previewed online"})
            return
        self._send_document(
            file_data, content_type, filename, download=download
        )

    def do_POST(self):
        if self.path == "/api/v1/login":
            self._login()
            return
        if self.path == "/api/v1/logout":
            self._logout()
            return
        if self.path == "/api/v1/documents":
            self._upload_document()
            return
        section = self.path.removeprefix("/api/v1/publish/")
        if section not in SECTIONS or self.path != f"/api/v1/publish/{section}":
            self._send_json(404, {"error": "Not found"})
            return
        publish_token = os.environ.get("CLEON_WEBSITE_PUBLISH_TOKEN", "")
        if not self._authorized(publish_token):
            self._send_json(401, {"error": "Unauthorized"})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            self._send_json(400, {"error": "Invalid Content-Length"})
            return
        if length <= 0 or length > MAX_BODY_BYTES:
            self._send_json(413, {"error": "Request body must be under 5 MB"})
            return
        try:
            payload = json.loads(self.rfile.read(length))
        except (json.JSONDecodeError, UnicodeDecodeError):
            self._send_json(400, {"error": "Request body must be valid JSON"})
            return
        if not isinstance(payload, dict) or not self._valid_payload(section, payload):
            self._send_json(400, {"error": f"Invalid {section} snapshot"})
            return
        try:
            published_at = datetime.fromisoformat(payload["published_at"])
            if published_at.tzinfo is None or published_at.utcoffset() is None:
                raise ValueError("published_at must include a timezone")
            published_at_value = published_at.astimezone(
                timezone.utc
            ).isoformat(timespec="microseconds")
            serialized_payload = json.dumps(
                payload, ensure_ascii=False, allow_nan=False
            )
        except (KeyError, TypeError, ValueError):
            self._send_json(
                400, {"error": "Snapshot contains invalid JSON values or timestamp"}
            )
            return

        updated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
        try:
            with _connect_database() as conn:
                cursor = conn.execute("""
                    INSERT INTO website_snapshots(
                        section,payload,updated_at,published_at
                    ) VALUES (?,?,?,?)
                    ON CONFLICT(section) DO UPDATE SET
                        payload=excluded.payload,
                        updated_at=excluded.updated_at,
                        published_at=excluded.published_at
                    WHERE excluded.published_at>=website_snapshots.published_at
                """, (
                    section, serialized_payload, updated_at,
                    published_at_value,
                ))
        except sqlite3.Error:
            self._send_json(500, {"error": "Could not store snapshot"})
            return
        if cursor.rowcount == 0:
            self._send_json(409, {"error": "A newer snapshot is already stored"})
            return
        self._send_json(200, {"section": section, "updated_at": updated_at})

    def do_GET(self):
        parsed_path = urlparse(self.path)
        if parsed_path.path not in (
            "/api/v1/me", "/api/v1/snapshot", "/api/v1/export",
            "/api/v1/documents",
        ) and not parsed_path.path.startswith("/api/v1/documents/"):
            self._send_json(404, {"error": "Not found"})
            return
        publish_token = os.environ.get("CLEON_WEBSITE_PUBLISH_TOKEN", "")
        is_publisher = self._authorized(publish_token)
        if parsed_path.path == "/api/v1/documents" and is_publisher:
            self._list_documents()
            return
        if (
            parsed_path.path.startswith("/api/v1/documents/")
            and (parse_qs(parsed_path.query).get("download") or [""])[0] == "1"
            and is_publisher
        ):
            self._get_document(parsed_path, {"role": "super_admin"})
            return
        try:
            member = self._current_member()
        except sqlite3.Error:
            self._send_json(500, {"error": "Could not verify member session"})
            return
        if member is None:
            self._send_json(401, {"error": "Sign in with a member account"})
            return
        if parsed_path.path == "/api/v1/me":
            self._send_json(200, {
                "email": member["email"],
                "role": member["role"],
                "expires_at": member["expires_at"],
            })
            return
        if parsed_path.path == "/api/v1/documents":
            self._list_documents()
            return
        if parsed_path.path.startswith("/api/v1/documents/"):
            self._get_document(parsed_path, member)
            return
        if parsed_path.path == "/api/v1/export":
            if member["role"] not in ("admin", "super_admin"):
                self._send_json(403, {"error": "Report downloads require an admin role"})
                return
            self._export_report(parse_qs(parsed_path.query))
            return
        try:
            with _connect_database() as conn:
                all_snapshots = {
                    section: {"data": json.loads(payload), "updated_at": updated_at}
                    for section, payload, updated_at in conn.execute(
                        "SELECT section,payload,updated_at FROM website_snapshots"
                    )
                }
        except (sqlite3.Error, json.JSONDecodeError):
            self._send_json(500, {"error": "Could not read snapshots"})
            return
        permitted_sections = ROLE_ACCESS[member["role"]]
        snapshots = {
            section: snapshot for section, snapshot in all_snapshots.items()
            if section in permitted_sections
        }
        self._send_json(200, snapshots)

    def _read_json_body(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
        except ValueError:
            self._send_json(400, {"error": "Invalid Content-Length"})
            return None
        if length <= 0 or length > 16 * 1024:
            self._send_json(413, {"error": "Request body is too large"})
            return None
        try:
            body = json.loads(self.rfile.read(length))
        except (json.JSONDecodeError, UnicodeDecodeError):
            self._send_json(400, {"error": "Request body must be valid JSON"})
            return None
        if not isinstance(body, dict):
            self._send_json(400, {"error": "Request body must be a JSON object"})
            return None
        return body

    def _login(self):
        address = self.client_address[0]
        now = time.monotonic()
        with LOGIN_ATTEMPTS_LOCK:
            attempts, window_start = LOGIN_ATTEMPTS.get(address, (0, now))
            if now - window_start >= 60:
                attempts, window_start = 0, now
            if attempts >= 10:
                self._send_json(429, {"error": "Too many sign-in attempts. Try again later."})
                return
            LOGIN_ATTEMPTS[address] = (attempts + 1, window_start)
        body = self._read_json_body()
        if body is None:
            return
        email = body.get("email")
        password = body.get("password")
        if not isinstance(email, str) or not isinstance(password, str):
            self._send_json(400, {"error": "Email and password are required"})
            return
        email = email.strip().casefold()
        if len(email) > 254 or len(password) > 1024:
            self._send_json(400, {"error": "Email or password is too long"})
            return
        try:
            with _connect_database() as conn:
                member = conn.execute(
                    "SELECT salt,password_hash,role FROM website_members WHERE email=?",
                    (email,),
                ).fetchone()
            salt = member[0] if member else "0" * 32
            password_hash = _password_hash(password, salt)
            valid = bool(
                member and hmac.compare_digest(password_hash, member[1])
                and member[2] in ROLES
            )
        except sqlite3.Error:
            self._send_json(500, {"error": "Could not verify member account"})
            return
        if not valid:
            self._send_json(401, {"error": "Invalid member email or password"})
            return
        token = secrets.token_urlsafe(32)
        expires_at = datetime.now(timezone.utc).timestamp() + SESSION_HOURS * 3600
        expires = datetime.fromtimestamp(
            expires_at, timezone.utc
        ).isoformat(timespec="seconds")
        try:
            with _connect_database() as conn:
                conn.execute(
                    "INSERT INTO website_member_sessions(token_hash,email,expires_at) "
                    "VALUES (?,?,?)",
                    (_token_hash(token), email, expires),
                )
                conn.execute(
                    "DELETE FROM website_member_sessions WHERE expires_at<=?",
                    (datetime.now(timezone.utc).isoformat(timespec="seconds"),),
                )
        except sqlite3.Error:
            self._send_json(500, {"error": "Could not create member session"})
            return
        with LOGIN_ATTEMPTS_LOCK:
            LOGIN_ATTEMPTS.pop(address, None)
        self._send_json(200, {
            "token": token,
            "email": email,
            "role": member[2],
            "expires_at": expires,
        })

    def _logout(self):
        token = self._bearer_token()
        if not token:
            self._send_json(401, {"error": "Sign in required"})
            return
        try:
            with _connect_database() as conn:
                conn.execute(
                    "DELETE FROM website_member_sessions WHERE token_hash=?",
                    (_token_hash(token),),
                )
        except sqlite3.Error:
            self._send_json(500, {"error": "Could not end member session"})
            return
        self._send_json(200, {"logged_out": True})

    def _current_member(self):
        token = self._bearer_token()
        if not token:
            return None
        now = datetime.now(timezone.utc).isoformat(timespec="seconds")
        with _connect_database() as conn:
            row = conn.execute(
                "SELECT members.email,members.role,sessions.expires_at "
                "FROM website_member_sessions AS sessions "
                "JOIN website_members AS members ON members.email=sessions.email "
                "WHERE sessions.token_hash=? AND sessions.expires_at>?",
                (_token_hash(token), now),
            ).fetchone()
        if row is None or row[1] not in ROLES:
            return None
        return {"email": row[0], "role": row[1], "expires_at": row[2]}

    def _bearer_token(self):
        scheme, separator, token = self.headers.get(
            "Authorization", ""
        ).partition(" ")
        if separator != " " or scheme.casefold() != "bearer":
            return ""
        return token

    def _export_report(self, query):
        report_name = (query.get("report") or [""])[0]
        export_format = (query.get("format") or [""])[0].casefold()
        if export_format not in ("csv", "xlsx", "pdf"):
            self._send_json(400, {"error": "Choose csv, xlsx, or pdf"})
            return
        try:
            with _connect_database() as conn:
                row = conn.execute(
                    "SELECT payload FROM website_snapshots WHERE section='reports'"
                ).fetchone()
            if not row:
                self._send_json(404, {"error": "No reports have been published"})
                return
            payload = json.loads(row[0])
            report = payload["reports"].get(report_name)
            if not report:
                self._send_json(404, {"error": "Report not found"})
                return
            headers = report["headers"]
            rows = report["rows"]
            safe_name = "".join(
                character.lower() if character.isalnum() else "_"
                for character in report_name
            ).strip("_")
            body, content_type, extension = self._format_export(
                export_format, report_name, headers, rows,
                payload["period"],
            )
        except (sqlite3.Error, json.JSONDecodeError, KeyError, TypeError) as error:
            self._send_json(500, {"error": f"Could not prepare report export: {error}"})
            return
        except ImportError as error:
            self._send_json(503, {"error": str(error)})
            return
        self._send_file(
            body, content_type, f"{safe_name}_report.{extension}"
        )

    @staticmethod
    def _format_export(export_format, report_name, headers, rows, period):
        def safe_string(value):
            text = "" if value is None else str(value)
            if text.startswith(("=", "+", "-", "@", "\t", "\r")):
                return "'" + text
            return text

        if export_format == "csv":
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow(headers)
            for row in rows:
                writer.writerow([
                    value if isinstance(value, (int, float)) else safe_string(value)
                    for value in row
                ])
            return (
                ("\ufeff" + output.getvalue()).encode("utf-8"),
                "text/csv; charset=utf-8",
                "csv",
            )

        if export_format == "xlsx":
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill
            except ImportError as error:
                raise ImportError("Excel export requires openpyxl.") from error
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = re.sub(r"[\[\]*?:/\\]", "_", report_name)[:31] or "Report"
            sheet.append([safe_string(header) for header in headers])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="24364B")
            for row in rows:
                sheet.append([
                    value if isinstance(value, (int, float)) else safe_string(value)
                    for value in row
                ])
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(
                    42,
                    max(12, max(len(str(cell.value or "")) for cell in column) + 2),
                )
                sheet.column_dimensions[column[0].column_letter].width = width
            output = io.BytesIO()
            workbook.save(output)
            return (
                output.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "xlsx",
            )

        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.platypus import (
                Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
            )
        except ImportError as error:
            raise ImportError("PDF export requires reportlab.") from error
        output = io.BytesIO()
        document = SimpleDocTemplate(
            output, pagesize=landscape(A4), title=f"CLEON {report_name} Report"
        )
        styles = getSampleStyleSheet()
        table_style = ParagraphStyle(
            "CLEONReportCell", parent=styles["BodyText"],
            fontSize=7, leading=8,
        )
        reportlab_text = lambda value: escape(
            "" if value is None else str(value)
        )
        content = [
            Paragraph(
                f"CLEON Oil &amp; Gas — {reportlab_text(report_name)}",
                styles["Title"],
            ),
            Paragraph(
                f"Period: {reportlab_text(period.get('from', ''))} to "
                f"{reportlab_text(period.get('to', ''))}",
                styles["Normal"],
            ),
            Spacer(1, 10),
        ]
        if headers:
            data = [[
                Paragraph(reportlab_text(value), table_style)
                for value in headers
            ]]
            data.extend(
                [
                    Paragraph(reportlab_text(value), table_style)
                    for value in row
                ]
                for row in rows
            )
            table = Table(
                data,
                colWidths=[
                    (landscape(A4)[0] - 144) / len(headers)
                ] * len(headers),
                repeatRows=1,
            )
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#24364B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [
                    colors.white, colors.HexColor("#F1F5F9")
                ]),
            ]))
            content.append(table)
        else:
            content.append(Paragraph("No report columns are available.", styles["Normal"]))
        document.build(content)
        return output.getvalue(), "application/pdf", "pdf"

    @staticmethod
    def _valid_payload(section, payload):
        if section == "dashboard":
            return (
                isinstance(payload.get("date"), str)
                and isinstance(payload.get("metrics"), dict)
                and isinstance(payload.get("published_at"), str)
            )
        reports = payload.get("reports")
        return (
            isinstance(payload.get("period"), dict)
            and isinstance(payload.get("summary"), dict)
            and isinstance(reports, dict)
            and isinstance(payload.get("published_at"), str)
            and all(
                isinstance(report, dict)
                and isinstance(report.get("headers"), list)
                and isinstance(report.get("rows"), list)
                for report in reports.values()
            )
        )

    def _authorized(self, expected_token):
        if not expected_token:
            return False
        scheme, separator, supplied_token = self.headers.get(
            "Authorization", ""
        ).partition(" ")
        return (
            separator == " "
            and scheme.casefold() == "bearer"
            and hmac.compare_digest(supplied_token, expected_token)
        )

    def log_message(self, format_string, *args):
        if self.command not in ("OPTIONS",):
            super().log_message(format_string, *args)


def main():
    publish_token = os.environ.get("CLEON_WEBSITE_PUBLISH_TOKEN", "")
    if len(publish_token) < 32:
        raise RuntimeError(
            "Set CLEON_WEBSITE_PUBLISH_TOKEN to a random value of at least "
            "32 characters."
        )
    _initialize_database()
    if len(sys.argv) > 1:
        if sys.argv[1] == "add-member" and len(sys.argv) in (3, 4):
            _add_member(
                sys.argv[2], sys.argv[3] if len(sys.argv) == 4 else "member"
            )
            return
        if len(sys.argv) == 3 and sys.argv[1] == "remove-member":
            _remove_member(sys.argv[2])
            return
        if len(sys.argv) == 4 and sys.argv[1] == "set-role":
            _set_member_role(sys.argv[2], sys.argv[3])
            return
        if len(sys.argv) == 2 and sys.argv[1] == "list-members":
            _list_members()
            return
        raise ValueError(
            "Usage: python website_api.py "
            "[add-member EMAIL [ROLE]|remove-member EMAIL|"
            "set-role EMAIL ROLE|list-members]"
        )
    host = os.environ.get("CLEON_WEBSITE_API_HOST", "127.0.0.1")
    port = int(os.environ.get(
        "CLEON_WEBSITE_API_PORT",
        os.environ.get("PORT", "8765"),
    ))
    server = ThreadingHTTPServer((host, port), WebsiteAPIHandler)
    print(f"CLEON website API listening on {host}:{port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("Stopping CLEON website API.")
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
